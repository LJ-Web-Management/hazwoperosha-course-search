"""Merge the B2B catalog workbook (HAZWOPER-OSHA_Master_Course_Catalog_and_B2B_Bundles_v3.xlsx)
into assets/data.js. See README.md ("Merging the B2B catalog workbook") for when to run it.

Usage:
    python3 tools/merge_b2b_catalog.py /path/to/workbook.xlsx [first_row last_row]

first_row/last_row are 1-based "Bundle Contents" sheet rows (default 43 74, the three Data
Center packs). Safe to re-run: courses/bundles already present are matched by name, not
duplicated.
"""
import json
import re
import sys

import openpyxl

DATA = "assets/data.js"
NAMES = ["COURSES", "BUNDLES", "BUNDLE_CONTENTS", "INDUSTRIES", "CATEGORIES"]
PACK_TYPE = "Compliance Pack"


def norm(s):
    return re.sub(r"[^a-z0-9]+", " ", (s or "").lower()).strip()


def fmt_duration(hours):
    if hours is None:
        return None
    if hours < 1:
        return f"{round(hours * 60)} Min"
    h = int(hours) if float(hours).is_integer() else hours
    return f"{h} Hour" if h == 1 else f"{h} Hours"


def split_citation(text):
    # "OSHA/NFPA – 29 CFR 1910 Subpart S / NFPA 70E" -> ("OSHA/NFPA", "29 CFR 1910 Subpart S / NFPA 70E")
    if not text:
        return None, None
    parts = text.split(" – ", 1)
    return (parts[0].strip(), parts[1].strip()) if len(parts) == 2 else (None, text.strip())


def main():
    src = sys.argv[1]
    first, last = (int(sys.argv[2]), int(sys.argv[3])) if len(sys.argv) > 3 else (43, 74)

    text = open(DATA, encoding="utf-8").read()
    header = text.splitlines()[0]
    data = {n: json.loads(re.search(r"^const %s = (.*);$" % n, text, re.M).group(1)) for n in NAMES}
    courses, bundles, contents = data["COURSES"], data["BUNDLES"], data["BUNDLE_CONTENTS"]
    industries = set(data["INDUSTRIES"])
    by_name = {norm(c["name"]): c for c in courses}

    wb = openpyxl.load_workbook(src, data_only=True)

    # Master Course List: add new courses; tag renamed ones with their new title.
    added, renamed = [], 0
    for r in wb["Master Course List"].iter_rows(min_row=2, values_only=True):
        if not r[1]:
            continue
        (_, name, citation, price, hours, _validity, _desc, category, _meta, orig_title,
         course_type, audience, _packs, libraries) = r[:14]
        existing = by_name.get(norm(name)) or (by_name.get(norm(orig_title)) if orig_title else None)
        if existing:
            # Workbook retitled an existing course: keep the old name, make the new one searchable.
            if name != existing["name"] and name not in existing["altTags"]:
                existing["altTags"].append(name)
                renamed += 1
            by_name[norm(name)] = existing
            continue
        reg_body, cit = split_citation(citation)
        libs = [l.strip() for l in (libraries or "").split(";") if l.strip()]
        tags = [t for t in (re.sub(r" Library$", " Bundle", l) for l in libs) if t in industries]
        course = {
            "id": max(c["id"] for c in courses) + 1, "category": category, "family": name,
            "name": name, "type": course_type, "regBody": reg_body, "citation": cit,
            "industries": audience, "duration": fmt_duration(hours), "msrp": price,
            "bundleClass": ("Certification-Tier (À La Carte)"
                            if "Train-the-Trainer" in (category or "") else "Bundle-Eligible"),
            "industryTags": tags, "altTags": [],
        }
        courses.append(course)
        by_name[norm(name)] = course
        added.append(name)

    # Bundle Contents rows first..last: build the packs and apply their per-seat pricing.
    pack_rows = {}
    ws = wb["Bundle Contents"]
    for r in ws.iter_rows(min_row=first, max_row=last, values_only=True):
        bname, cname, category, hours, price = r[:5]
        pack_rows.setdefault(bname, []).append((cname, category, hours, price))

    overview = {r[0]: r for r in wb["B2B Bundles & Seat Tiers"].iter_rows(min_row=2, values_only=True) if r[0]}
    repriced = []
    for bname, rows in pack_rows.items():
        o = overview[bname]
        (_, model, buyer, what, n_courses, _total_hours, ala_value, list_price, discount) = o[:9]
        bid = f"{PACK_TYPE}||{bname}"
        items = []
        for cname, category, hours, price in rows:
            course = by_name[norm(cname)]
            if course["msrp"] != price:
                repriced.append((course["name"], course["msrp"], price))
                course["msrp"] = price
            items.append({
                "name": course["name"], "category": course["category"], "type": course["type"],
                "duration": course["duration"], "msrp": price,
                "status": "Included in Flat Bundle Price",
            })
        cats = sorted({i["category"] for i in items})
        bundle = {
            "id": bid, "type": PACK_TYPE, "name": bname, "scope": ", ".join(cats),
            "totalCourses": n_courses, "inclFlat": len(items), "alaCarteAddons": 0,
            "costSeparate": ala_value, "priceTier": model, "bundlePrice": list_price,
            "savings": discount,
            "alaCarteTerms": f"All courses included in flat pack price. Target buyer: {buyer}. {what}",
        }
        bundles[:] = [b for b in bundles if b["id"] != bid] + [bundle]
        contents[bid] = items

    # Keep every bundle's contents in step with the new course prices.
    new_price = {n: p for n, _, p in repriced}
    for items in contents.values():
        for i in items:
            if i["name"] in new_price:
                i["msrp"] = new_price[i["name"]]

    data["CATEGORIES"] = sorted({c["category"] for c in courses})
    with open(DATA, "w", encoding="utf-8") as f:
        f.write(header + "\n")
        for n in NAMES:
            f.write(f"const {n} = " + json.dumps(data[n], ensure_ascii=False) + ";\n")

    print(f"added {len(added)} courses, tagged {renamed} renamed courses, "
          f"repriced {len(repriced)}, packs: {list(pack_rows)}")
    for n, old, new in repriced:
        print(f"  repriced {n}: {old} -> {new}")


if __name__ == "__main__":
    main()
